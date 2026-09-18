import asyncio
import os
from aiogram import Bot, Dispatcher, F
from aiogram.types import Message
import openpyxl
import requests
import json
from bs4 import BeautifulSoup
from aiogram.filters import Command

with open('config.json', 'r', encoding='utf-8') as f:
    config = json.load(f)

bot = Bot(token=config['tg_token_done'])
dp = Dispatcher()


@dp.message()
async def alo(message: Message):
    if "done" in message.text:
        text = await delit_akks()
        text = text.split('|')
        text.pop(len(text)-1)
        await message.answer("почта:пароль_почты:логин:пароль:ссылка")
        for k in range(len(text)):
            await message.answer(text[k])
    elif "ерекличка" in message.text:
        await message.answer("я не сплю")
    elif "колько" in message.text:
        count = await search()
        await message.answer(str(count))
    elif "апас" in message.text:
        count = await search_all()
        await message.answer(str(count))

async def main():
    await dp.start_polling(bot)

async def delit_akks():
    akks_done = ''
    book = openpyxl.load_workbook("done.xlsx")
    sheet = book.active
    nomer_stroki=0
    k=0
    responce = requests.get("http://localhost:1242/Api/Bot/ASF")
    data = responce.json()["Result"]
    while sheet[nomer_stroki+1][0].value != None:
        if int(sheet[nomer_stroki+1][1].value) > 9:
            url = "https://steamcommunity.com/profiles/" + data[sheet[nomer_stroki+1][0].value]["s_SteamID"]
            akks_done += f"{sheet[nomer_stroki + 1][3].value}:{sheet[nomer_stroki + 1][4].value}:{sheet[nomer_stroki + 1][0].value}:{sheet[nomer_stroki + 1][2].value}:{url}|"
            os.remove(f"asf/config/{sheet[nomer_stroki+1][0].value}.json")
            os.remove(f"asf/config/{sheet[nomer_stroki+1][0].value}.db")
            sheet.delete_rows(nomer_stroki+1)
            k += 1
        nomer_stroki += 1
    book.save("done.xlsx")
    book.close()
    return akks_done

async def search():
    book = openpyxl.load_workbook("done.xlsx")
    sheet = book.active
    nomer_stroki=0
    k=0
    while sheet[nomer_stroki+1][0].value != None:
        if int(sheet[nomer_stroki+1][1].value) > 9:
            k+=1
        nomer_stroki+=1
    book.save("done.xlsx")
    book.close()
    return k

async def search_all():
    book = openpyxl.load_workbook("done.xlsx")
    sheet = book.active
    nomer_stroki=0
    k=0
    while sheet[nomer_stroki+1][0].value != None:
        nomer_stroki+=1
    book.save("done.xlsx")
    book.close()
    return nomer_stroki



asyncio.run(main())