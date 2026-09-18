import requests
import asyncio
import openpyxl
from ASF import IPC
import time
import json

async def command(asf, cmd):
    return await asf.Api.Command.post(body={
        'Command': cmd
    })

def key_reborn(text):
    key = [['', '', '', '', '', '', '', '', '', ''],
           ['', '', '', '', '', '', '', '', '', ''],
           ['', '', '', '', '', '', '', '', '', ''],
           ['', '', '', '', '', '', '', '', '', '']]
    keys = ['', '', '', '', '']
    text = text.split("Вместо * цифра от 0 до 9")
    text = text[1].split('\n')
    keys[1] = text[1].replace("🔑 ", "")
    keys[2] = text[2].replace("🔑 ", "")
    keys[3] = text[3].replace("🔑 ", "")
    keys[4] = text[4].replace("🔑 ", "")
    for k in range(10):
        key[0][k] = keys[1].replace("*", str(k))
    for k in range(10):
        key[1][k] = keys[2].replace("*", str(k))
    for k in range(10):
        key[2][k] = keys[3].replace("*", str(k))
    for k in range(10):
        key[3][k] = keys[4].replace("*", str(k))
    return key

def telega(msg):
    token_tg = config["tg_token_main"]
    chat_id = config["chat_id"]
    url_req = "https://api.telegram.org/bot" + token_tg + "/sendMessage" + "?chat_id=" + chat_id + "&text=" + msg
    results = requests.get(url_req)


async def main():
    async with IPC(ipc='http://127.0.0.1:1242', password='YOUR IPC PASSWORD') as asf:

        i = 0
        token = config["token"]
        token_prof = config["token_prof"]
        version = 5.131
        domain = config["domain_group"]
        idgruppy = config["id_group"]

        respons = requests.get("https://api.vk.com/method/wall.get",
                               params={
                                   "access_token": token,
                                   "v": version,
                                   "domain": domain
                               }
                               )

        data = respons.json()["response"]["items"][0o0]["id"]
        proverka = data

        print("запущенно")
        dada = 0
        """while dada != 1:
            dada = 1
            if 1 == 2:
                print("1")"""

        while True:
            if proverka == data:
                respons = requests.get("https://api.vk.com/method/wall.get",
                                       params={
                                           "access_token": token,
                                           "v": version,
                                           "domain": domain
                                       }
                                       )
                data = respons.json()["response"]["items"][0o0]["id"]
            else:
                timeStart = time.time()
                proverka = data
                respons = requests.get("https://api.vk.com/method/wall.get",
                                       params={
                                           "access_token": token,
                                           "v": version,
                                           "domain": domain
                                       }
                                       )
                text = respons.json()["response"]["items"][0o0]["text"]

                keys = key_reborn(text)


                akk = 1
                for l in range(10):
                    book = openpyxl.load_workbook("done.xlsx")
                    sheet = book.active
                    while int(sheet[akk][1].value) > 10:
                        akk +=1
                    book.save("done.xlsx")
                    book.close()
                    cmd =str(sheet[akk][0].value)
                    url = f"http://localhost:1242/Api/Bot/{cmd}/Redeem"
                    kluchi = ""
                    for c in range(len(keys)):
                        kluchi= kluchi +f'{keys[c][l]}|'
                    kluchi = kluchi.split("|")
                    kluchi.pop(len(kluchi) - 1)
                    print(kluchi)
                    respons = requests.post(url, json={"KeysToRedeem":kluchi})
                    resp = respons.json()['Result'][sheet[akk][0].value]
                    delit = 0
                    for c in range(len(kluchi)):
                        print(kluchi[c] + " = " + str(resp[kluchi[c]]['purchase_result_details']))
                        if resp[kluchi[c]]['purchase_result_details'] == 15:
                            keys.pop(c-delit)
                            delit+=1
                            print("кем-то актив")
                        if resp[kluchi[c]]['purchase_result_details'] == 9:
                            for k in range(10):
                                keys[c-delit][k] = keys[c-delit][l]
                            print("повторка")
                        if resp[kluchi[c]]['purchase_result_details'] == 0:
                            book = openpyxl.load_workbook("done.xlsx")
                            sheet = book.active
                            x = int(sheet[akk][1].value)
                            x += 1
                            sheet[akk][1].value = x
                            book.save("done.xlsx")
                            keys.pop(c-delit)
                            delit+=1
                            print("ок")

                    """if "OK" in resp.result:
                        book = openpyxl.load_workbook("done.xlsx")
                        sheet = book.active
                        x = int(sheet[akk][1].value)
                        x += 1
                        sheet[akk][1].value = x
                        book.save("done.xlsx")
                        j=0
                        while j < len(delit):
                            if "OK" in delit[j]:
                                keys.pop(j)
                                delit.pop(j)
                                j-=1
                            j+=1
                    if "DuplicateActivationCode" in resp.result:
                        j = 0
                        while j < len(delit):
                            if "DuplicateActivationCode" in delit[j]:
                                keys.pop(j)
                                delit.pop(j)
                                j -= 1
                            j += 1
                    if "AlreadyPurchased" in resp.result:
                        j = 0
                        while j < len(delit):
                            if "AlreadyPurchased" in delit[j]:
                                for k in range(10):
                                    keys[j][k]=keys[j][l]
                            j += 1"""

                    akk +=1
                    print("конец партии")
                timeEnd = time.time()
                print(timeEnd)
                print(timeStart)
                print(timeEnd-timeStart)




while True:

    with open('config.json', 'r', encoding='utf-8') as f:
        config = json.load(f)
    try:
        loop = asyncio.get_event_loop()
        output = loop.run_until_complete(main())
        loop.close()
    except Exception as e:
        print(e)
        telega("ошибка")
